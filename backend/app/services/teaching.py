"""Document extraction and deterministic, holiday-aware lecture scheduling."""
from datetime import date, datetime, time, timedelta, timezone
from io import BytesIO
from pathlib import Path
import re
import zipfile
from pydantic import BaseModel, ConfigDict, Field, field_validator, model_validator

class Topic(BaseModel):
    model_config = ConfigDict(extra="forbid", str_strip_whitespace=True)
    title: str = Field(min_length=1, max_length=180)
    lectures: int = Field(default=1, ge=1, le=40)
    preparation: str = Field(default="", max_length=1000)

class WeeklySlot(BaseModel):
    weekday: int = Field(ge=0, le=6)
    starts: time
    minutes: int = Field(default=60, ge=15, le=240)

    @field_validator("starts")
    @classmethod
    def local_time(cls, value):
        if value.tzinfo or value.second or value.microsecond:
            raise ValueError("Use a local time in HH:MM format.")
        if value.hour * 60 + value.minute > 1440 - 15:
            raise ValueError("Class time is too late in the day.")
        return value

class PlanInput(BaseModel):
    model_config = ConfigDict(extra="forbid", str_strip_whitespace=True)
    course: str = Field(min_length=1, max_length=120)
    assignee: str = Field(min_length=1, max_length=120)
    start_date: date
    end_date: date
    topics: list[Topic] = Field(min_length=1, max_length=200)
    slots: list[WeeklySlot] = Field(min_length=1, max_length=21)
    holidays: list[date] = Field(default_factory=list, max_length=366)
    reminder_minutes: int = Field(default=15, ge=5, le=1440)

    @model_validator(mode="after")
    def validate_schedule(self):
        if not 0 <= (self.end_date - self.start_date).days <= 366:
            raise ValueError("Semester dates must be ordered and span at most one year.")
        for slot in self.slots:
            if slot.starts.hour * 60 + slot.starts.minute + slot.minutes > 1440:
                raise ValueError("Classes must finish on the same day.")
        for index, slot in enumerate(self.slots):
            start = slot.starts.hour * 60 + slot.starts.minute
            for other in self.slots[:index]:
                other_start = other.starts.hour * 60 + other.starts.minute
                if slot.weekday == other.weekday and start < other_start + other.minutes and other_start < start + slot.minutes:
                    raise ValueError("Weekly class slots overlap.")
        return self

def extract_document(filename: str, content: bytes) -> str:
    if not content or len(content) > 5 * 1024 * 1024:
        raise ValueError("Upload a nonempty document of at most 5 MB.")
    suffix = Path(filename).suffix.lower()
    if suffix in {".txt", ".md", ".csv"}:
        text = content.decode("utf-8-sig")
    elif suffix == ".pdf":
        from pypdf import PdfReader
        if not content.startswith(b"%PDF-"):
            raise ValueError("The file is not a PDF document.")
        reader = PdfReader(BytesIO(content))
        if reader.is_encrypted or len(reader.pages) > 100:
            raise ValueError("Use an unencrypted PDF with at most 100 pages.")
        parts = []
        for page in reader.pages:
            stream = page.get_contents()
            if stream and len(stream.get_data()) > 5 * 1024 * 1024:
                raise ValueError("A PDF page is too complex; export it as text.")
            parts.append(page.extract_text() or "")
            if sum(map(len, parts)) > 60000:
                raise ValueError("Document text exceeds 60,000 characters.")
        text = "\n".join(parts)
    elif suffix == ".docx":
        from docx import Document
        with zipfile.ZipFile(BytesIO(content)) as archive:
            if sum(info.file_size for info in archive.infolist()) > 20 * 1024 * 1024:
                raise ValueError("Expanded Word document exceeds 20 MB.")
        doc = Document(BytesIO(content))
        text = "\n".join([p.text for p in doc.paragraphs] + [" | ".join(c.text for c in row.cells) for table in doc.tables for row in table.rows])
    else:
        raise ValueError("Supported documents: PDF, DOCX, TXT, Markdown and CSV.")
    text = text.replace("\x00", "").strip()
    if not text:
        raise ValueError("No readable text found. For a scanned PDF, use OCR first or paste the syllabus text.")
    if len(text) > 60000:
        raise ValueError("Document text exceeds 60,000 characters.")
    return text

def holiday_dates(text: str) -> list[str]:
    dates = set()
    for value in re.findall(r"\b\d{4}-\d{2}-\d{2}\b", text):
        try:
            dates.add(date.fromisoformat(value).isoformat())
        except ValueError:
            pass
    return sorted(dates)

def schedule_plan(payload: PlanInput) -> dict:
    # College workspace uses Asia/Kolkata: fixed UTC+05:30, no DST.
    zone = timezone(timedelta(hours=5, minutes=30))
    slots = []
    current = payload.start_date
    holidays = set(payload.holidays)
    while current <= payload.end_date:
        if current not in holidays:
            for slot in sorted(payload.slots, key=lambda s: s.starts):
                if current.weekday() == slot.weekday:
                    start = datetime.combine(current, slot.starts, zone)
                    slots.append((start, start + timedelta(minutes=slot.minutes)))
        current += timedelta(days=1)
    required = sum(t.lectures for t in payload.topics)
    sessions = []
    index = 0
    for topic in payload.topics:
        for part in range(topic.lectures):
            if index >= len(slots):
                break
            start, end = slots[index]
            sessions.append({"title": f"{payload.course}: {topic.title}" + (f" ({part + 1}/{topic.lectures})" if topic.lectures > 1 else ""),
                "instructions": f"Teach: {topic.title}.\n" + topic.preparation,
                "starts_at": start.astimezone(timezone.utc).isoformat(), "ends_at": end.astimezone(timezone.utc).isoformat(),
                "assignee": payload.assignee, "blocks_asset": False, "asset_id": None})
            index += 1
    return {"sessions": sessions, "required_lectures": required, "available_slots": len(slots),
            "unallocated_lectures": max(0, required - len(slots)), "spare_slots": max(0, len(slots) - required),
            "timezone": "Asia/Kolkata"}
